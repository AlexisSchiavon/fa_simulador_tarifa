"""Generación de reportes Excel profesionales para escenarios Hyperred FA."""
from __future__ import annotations

import io
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Paleta Hyperred FA
COLOR_HEADER_BG = "1F2937"
COLOR_HEADER_FG = "FFFFFF"
COLOR_FA_AMARILLO = "FFB800"
COLOR_CLASE_E = "378ADD"
COLOR_CLASE_S = "639922"
COLOR_CLASE_P = "BA7517"
COLOR_CLASE_X = "D85A30"
COLOR_ALTERNATING = "F9FAFB"

CLASE_COLORS = {
    "E": COLOR_CLASE_E,
    "S": COLOR_CLASE_S,
    "P": COLOR_CLASE_P,
    "X": COLOR_CLASE_X,
}

FMT_MXN = '"$"#,##0.00'
FMT_PCT = "0.0%"


def _header_style(cell, bg_color: str = COLOR_HEADER_BG, fg_color: str = COLOR_HEADER_FG):
    cell.font = Font(bold=True, color=fg_color, size=10)
    cell.fill = PatternFill("solid", fgColor=bg_color)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _set_col_widths(ws, widths: list[int]):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def generar_excel_escenario(
    escenario: dict,
    boletos: list[dict],
    servicios_catalogo: list[dict],
    comparacion: dict | None = None,
) -> bytes:
    """Genera un .xlsx completo con múltiples hojas para un escenario.

    Args:
        escenario: dict con metadatos del escenario.
        boletos: lista de boletos vendidos en el escenario.
        servicios_catalogo: catálogo de servicios adicionales.
        comparacion: resultados del comparador de 3 modelos (opcional).

    Returns:
        Bytes del archivo .xlsx.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # quitar hoja default

    _hoja_resumen(wb, escenario, boletos)
    _hoja_transacciones(wb, boletos)
    _hoja_evolucion_precio(wb, boletos)
    _hoja_servicios(wb, boletos, servicios_catalogo)
    if comparacion:
        _hoja_comparacion(wb, comparacion)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def _hoja_resumen(wb: openpyxl.Workbook, escenario: dict, boletos: list[dict]):
    ws = wb.create_sheet("Resumen")

    # Título
    ws.merge_cells("A1:D1")
    titulo = ws["A1"]
    titulo.value = "Hyperred FA Sandbox — Resumen del Escenario"
    titulo.font = Font(bold=True, size=14, color=COLOR_HEADER_BG)
    titulo.alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:D2")
    subtitulo = ws["A2"]
    subtitulo.value = "GENIE S.C. × Lumixia · Proyecto Hyperred 25061 · 2026"
    subtitulo.font = Font(italic=True, size=9, color="6B7280")
    subtitulo.alignment = Alignment(horizontal="center")

    ws.row_dimensions[3].height = 6

    datos = [
        ("Escenario", escenario.get("nombre", "—")),
        ("Modelo", escenario.get("modelo_principal", "hibrido").upper()),
        ("Ruta", escenario.get("ruta_id", "—")),
        ("Fecha generación", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("", ""),
        ("Boletos vendidos", len(boletos)),
        ("Ingreso boletos (MXN)", sum(b.get("precio_pagado_mxn", 0) for b in boletos)),
        ("Ingreso servicios (MXN)", sum(
            sum(s.get("precio_pagado_mxn", 0) for s in b.get("servicios", []))
            for b in boletos
        )),
        ("Precio promedio (MXN)", (
            sum(b.get("precio_pagado_mxn", 0) for b in boletos) / len(boletos)
            if boletos else 0
        )),
        ("Precio mínimo (MXN)", min((b.get("precio_pagado_mxn", 0) for b in boletos), default=0)),
        ("Precio máximo (MXN)", max((b.get("precio_pagado_mxn", 0) for b in boletos), default=0)),
    ]

    for row_idx, (label, value) in enumerate(datos, 4):
        cell_lbl = ws.cell(row=row_idx, column=1, value=label)
        cell_val = ws.cell(row=row_idx, column=2, value=value)
        if label:
            cell_lbl.font = Font(bold=True, size=10)
        if isinstance(value, float):
            cell_val.number_format = FMT_MXN
            cell_val.font = Font(bold=True, size=11, color=COLOR_HEADER_BG)

    _set_col_widths(ws, [30, 25, 20, 20])


def _hoja_transacciones(wb: openpyxl.Workbook, boletos: list[dict]):
    ws = wb.create_sheet("Transacciones")

    headers = ["#", "Clase", "Segmento", "Días Ant.", "Precio (MXN)",
               "Ing. Servicios (MXN)", "Total (MXN)", "Modelo", "Estado"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        _header_style(c)

    ws.freeze_panes = "A2"

    for row_idx, b in enumerate(boletos, 2):
        clase = b.get("clase_tarifaria", "S")
        precio = float(b.get("precio_pagado_mxn", 0))
        ing_srv = sum(float(s.get("precio_pagado_mxn", 0)) for s in b.get("servicios", []))

        ws.cell(row=row_idx, column=1, value=row_idx - 1)
        clase_cell = ws.cell(row=row_idx, column=2, value=clase)
        clase_cell.font = Font(bold=True, color=CLASE_COLORS.get(clase, "000000"))
        ws.cell(row=row_idx, column=3, value=b.get("segmento_pasajero", "ocio"))
        ws.cell(row=row_idx, column=4, value=b.get("dias_anticipacion", 0))
        precio_cell = ws.cell(row=row_idx, column=5, value=precio)
        precio_cell.number_format = FMT_MXN
        srv_cell = ws.cell(row=row_idx, column=6, value=ing_srv)
        srv_cell.number_format = FMT_MXN
        total_cell = ws.cell(row=row_idx, column=7, value=precio + ing_srv)
        total_cell.number_format = FMT_MXN
        total_cell.font = Font(bold=True)
        ws.cell(row=row_idx, column=8, value=b.get("modelo_usado", "hibrido"))
        ws.cell(row=row_idx, column=9, value=b.get("estado", "activo"))

        if row_idx % 2 == 0:
            for col in range(1, 10):
                ws.cell(row=row_idx, column=col).fill = PatternFill("solid", fgColor=COLOR_ALTERNATING)

    # Fila de totales
    total_row = len(boletos) + 2
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    total_precio = ws.cell(row=total_row, column=5,
                           value=sum(float(b.get("precio_pagado_mxn", 0)) for b in boletos))
    total_precio.number_format = FMT_MXN
    total_precio.font = Font(bold=True, color=COLOR_HEADER_BG)
    total_srv = ws.cell(row=total_row, column=6,
                        value=sum(
                            sum(float(s.get("precio_pagado_mxn", 0)) for s in b.get("servicios", []))
                            for b in boletos
                        ))
    total_srv.number_format = FMT_MXN
    total_srv.font = Font(bold=True, color=COLOR_HEADER_BG)
    total_t = ws.cell(row=total_row, column=7,
                      value=sum(
                          float(b.get("precio_pagado_mxn", 0)) +
                          sum(float(s.get("precio_pagado_mxn", 0)) for s in b.get("servicios", []))
                          for b in boletos
                      ))
    total_t.number_format = FMT_MXN
    total_t.font = Font(bold=True, size=12, color=COLOR_HEADER_BG)

    _set_col_widths(ws, [5, 8, 12, 10, 18, 20, 18, 12, 10])


def _hoja_evolucion_precio(wb: openpyxl.Workbook, boletos: list[dict]):
    ws = wb.create_sheet("Evolución Precio")

    headers = ["Boleto #", "Precio (MXN)", "Multiplicador", "Clase",
               "f_ocu", "f_ant", "f_dem", "f_seg", "f_lix"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        _header_style(c)

    ws.freeze_panes = "A2"

    for row_idx, b in enumerate(boletos, 2):
        factores = b.get("factores_desglose", {})
        ws.cell(row=row_idx, column=1, value=row_idx - 1)
        p = ws.cell(row=row_idx, column=2, value=float(b.get("precio_pagado_mxn", 0)))
        p.number_format = FMT_MXN
        ws.cell(row=row_idx, column=3, value=float(b.get("multiplicador_aplicado", 1)))
        clase = b.get("clase_tarifaria", "S")
        cc = ws.cell(row=row_idx, column=4, value=clase)
        cc.font = Font(bold=True, color=CLASE_COLORS.get(clase, "000000"))
        for fi, key in enumerate(["f_ocu", "f_ant", "f_dem", "f_seg", "f_lix"], 5):
            val = factores.get(key, 1.0)
            cell = ws.cell(row=row_idx, column=fi, value=float(val))
            if float(val) > 1.0:
                cell.font = Font(color="D85A30")
            elif float(val) < 1.0:
                cell.font = Font(color="639922")

    _set_col_widths(ws, [10, 16, 14, 8, 10, 10, 10, 10, 10])


def _hoja_servicios(wb: openpyxl.Workbook, boletos: list[dict], servicios_catalogo: list[dict]):
    ws = wb.create_sheet("Servicios Adicionales")

    headers = ["Boleto #", "Servicio ID", "Nombre", "Tier", "Precio (MXN)"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        _header_style(c)

    nombres = {s["id"]: s["nombre"] for s in servicios_catalogo}

    row_idx = 2
    for b_num, b in enumerate(boletos, 1):
        for srv in b.get("servicios", []):
            ws.cell(row=row_idx, column=1, value=b_num)
            ws.cell(row=row_idx, column=2, value=srv.get("servicio_id", ""))
            ws.cell(row=row_idx, column=3, value=nombres.get(srv.get("servicio_id", ""), ""))
            ws.cell(row=row_idx, column=4, value=srv.get("tier", "base"))
            p = ws.cell(row=row_idx, column=5, value=float(srv.get("precio_pagado_mxn", 0)))
            p.number_format = FMT_MXN
            row_idx += 1

    if row_idx == 2:
        ws.cell(row=2, column=1, value="(Sin servicios adicionales registrados)")

    _set_col_widths(ws, [10, 12, 35, 10, 16])


def _hoja_comparacion(wb: openpyxl.Workbook, comparacion: dict):
    ws = wb.create_sheet("Comparación Modelos")

    headers = ["Métrica", "Aviación Puro", "Flixbus Puro", "Híbrido FA"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        bg = COLOR_FA_AMARILLO if col == 4 else COLOR_HEADER_BG
        fg = COLOR_HEADER_BG if col == 4 else COLOR_HEADER_FG
        _header_style(c, bg_color=bg, fg_color=fg)

    metricas = [
        ("Ingreso boletos (MXN)", "ingreso_boletos", FMT_MXN),
        ("Ingreso servicios (MXN)", "ingreso_servicios", FMT_MXN),
        ("INGRESO TOTAL (MXN)", "_total", FMT_MXN),
        ("Precio promedio (MXN)", "precio_promedio", FMT_MXN),
        ("Precio mínimo (MXN)", "precio_minimo", FMT_MXN),
        ("Precio máximo (MXN)", "precio_maximo", FMT_MXN),
        ("Boletos simulados", "num_boletos", None),
    ]

    for row_idx, (label, key, fmt) in enumerate(metricas, 2):
        ws.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
        for col, modelo in enumerate(["aviacion", "flixbus", "hibrido"], 2):
            m = comparacion.get(modelo, {})
            if key == "_total":
                val = m.get("ingreso_boletos", 0) + m.get("ingreso_servicios", 0)
            else:
                val = m.get(key, 0)
            cell = ws.cell(row=row_idx, column=col, value=val)
            if fmt:
                cell.number_format = fmt
            if modelo == "hibrido":
                cell.font = Font(bold=True, color=COLOR_HEADER_BG)

    _set_col_widths(ws, [28, 18, 18, 18])
